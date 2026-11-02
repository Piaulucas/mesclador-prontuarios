"""Leitura da auditoria do mês, sem alterar planilhas ou pastas de prontuários."""

import hashlib
import json
import os
import re
import tempfile
import unicodedata
from collections import Counter
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string

from .config import MESES
from . import planilha, pdfs


AMARELOS = {"FFFFFF00", "FFFFFF99", "FFFFE699"}
VERMELHOS = {"FFFF0000", "FFFFC7CE", "FFFF9999"}
PADRAO_PDF = re.compile(r"^(\d+)\s*-\s*(.+)\.pdf$", re.IGNORECASE)
FAIXAS_EQUIPE = (6, 10)


def caminho_registros(config, empresa, mes):
    """Um arquivo por empresa e mês, sem usar o nome da empresa como caminho."""
    identificador = hashlib.sha256(empresa.encode("utf-8")).hexdigest()[:16]
    return Path(config.pasta_registros).expanduser() / f"{mes}-{identificador}.json"


def ler_registros(config, empresa, mes):
    caminho = caminho_registros(config, empresa, mes)
    if not caminho.exists():
        return {"fichas": {}}
    with caminho.open(encoding="utf-8") as arquivo:
        return json.load(arquivo)


def salvar_registros(config, empresa, mes, registros):
    """Substitui o JSON completo somente após a gravação do temporário."""
    caminho = caminho_registros(config, empresa, mes)
    caminho.parent.mkdir(parents=True, exist_ok=True)
    temporario = None
    try:
        with tempfile.NamedTemporaryFile("w", encoding="utf-8", dir=caminho.parent,
                                         prefix=".registro-", suffix=".tmp", delete=False) as arquivo:
            temporario = arquivo.name
            json.dump(registros, arquivo, ensure_ascii=False, indent=2)
            arquivo.flush()
            os.fsync(arquivo.fileno())
        os.replace(temporario, caminho)
    finally:
        if temporario and os.path.exists(temporario):
            os.unlink(temporario)


def aplicar_registros(fichas, registros, hoje):
    """Associa registros conferidos e calcula dias a partir do último envio."""
    resultado, inconsistencias = [], []
    for ficha in fichas:
        item = dict(ficha)
        registro = registros.get("fichas", {}).get(ficha["ficha"])
        if registro and not planilha.nomes_conferem(registro.get("paciente", ""), ficha["paciente"]):
            inconsistencias.append({"ficha": ficha["ficha"], "tipo": ficha["tipo"],
                                    "motivo": "registro_outro_paciente"})
            registro = None
        if registro:
            item["correcao"] = registro.get("correcao")
            envios = registro.get("envios", [])
            if envios:
                ultimo = max(enumerate(envios), key=lambda par: (par[1]["data"], par[0]))[1]
                item["ultimo_envio"] = ultimo
                dias = (hoje - date.fromisoformat(ultimo["data"])).days
                item["dias"] = dias
                item["faixa"] = ("amarelo" if dias <= FAIXAS_EQUIPE[0] else
                                 "laranja" if dias <= FAIXAS_EQUIPE[1] else "vermelho")
                item["devolucao"] = f"Enviada em {ultimo['data']} · {dias} dia(s) com a equipe"
        if "ultimo_envio" not in item:
            item["devolucao"] = "devolução não registrada"
        resultado.append(item)
    return resultado, inconsistencias


def pasta_prontuarios(config, empresa, ano, mes):
    """Resolve a pasta do mês; aceita MARCO quando o molde MARÇO não existe."""
    molde = config.auditoria_prontuarios.get(empresa)
    if not molde:
        return None
    caminho = Path(molde.format(ano=ano, mm=f"{int(mes):02d}", MES=MESES[int(mes) - 1])).expanduser()
    if caminho.is_dir():
        return caminho
    mes_sem_acento = "".join(c for c in unicodedata.normalize("NFKD", MESES[int(mes) - 1])
                             if not unicodedata.combining(c))
    variante = Path(molde.format(ano=ano, mm=f"{int(mes):02d}", MES=mes_sem_acento)).expanduser()
    return variante if variante.is_dir() else None


def localizar_prontuarios(config, empresa, ano, mes):
    """Devolve a pasta e um aviso quando ela não estiver disponível."""
    pasta = pasta_prontuarios(config, empresa, ano, mes)
    return {"pasta": pasta,
            "aviso": None if pasta else "Pasta de prontuários do mês não encontrada."}


def _pasta_do_mes(nome, ano, mes):
    """Reconhece mês e rejeita ano explícito diferente do consultado."""
    partes = re.findall(r"[A-Z]+|\d+", planilha.normalizar(nome))
    numero = f"{mes:02d}"
    if not (numero in partes or str(mes) in partes or
            planilha.normalizar(MESES[mes - 1]) in partes):
        return False
    anos = [parte for parte in partes if len(parte) == 4 and parte.startswith("20")]
    anos += [parte for parte in partes if parte.isdigit() and len(parte) == 2 and parte not in
             {f"{n:02d}" for n in range(1, 13)}]
    return not anos or all(parte in {str(ano), f"{ano % 100:02d}"} for parte in anos)


def localizar_scanners(config, empresa, ano, mes):
    """Localiza o atalho e, quando necessário, a pasta do mês mais recente."""
    configuracao = config.auditoria_scanners.get(empresa)
    if not configuracao:
        return {"pasta": None, "formato": None, "aviso": "Pasta de scanners não configurada."}
    formato = configuracao.get("formato")
    if formato not in {"pastas_por_paciente", "arquivos_soltos"}:
        return {"pasta": None, "formato": formato, "aviso": "Formato de scanners inválido."}
    raiz = Path(configuracao.get("pasta") or "").expanduser()
    if not configuracao.get("pasta") or not raiz.is_dir():
        return {"pasta": None, "formato": formato, "aviso": "Pasta de scanners não encontrada."}
    if formato == "arquivos_soltos":
        return {"pasta": raiz, "formato": formato, "aviso": None}
    try:
        with os.scandir(raiz) as entradas:
            pastas = [(entrada.stat().st_mtime, entrada.name, Path(entrada.path))
                      for entrada in entradas if entrada.is_dir() and _pasta_do_mes(entrada.name, ano, mes)]
    except OSError as erro:
        return {"pasta": None, "formato": formato, "aviso": f"Não foi possível listar scanners: {erro}"}
    if not pastas:
        return {"pasta": None, "formato": formato, "aviso": "Pasta de scanners do mês não encontrada."}
    pastas.sort(reverse=True)
    return {"pasta": pastas[0][2], "formato": formato,
            "aviso": "Mais de uma pasta de scanners do mês; usada a mais recente." if len(pastas) > 1 else None}


def listar_scanners(local, ano, mes, ilegiveis=None):
    """Lista somente nomes e metadados; nunca abre o conteúdo escaneado.

    Uma subpasta que o sistema não consegue abrir (o Google Drive às vezes responde
    "Resource deadlock avoided" para uma pasta que ainda está baixando) não derruba a
    leitura da empresa: o nome dela vai para `ilegiveis`, se a lista for passada, e a
    pasta do paciente continua contando como escaneada, com a data da própria pasta.
    """
    pasta, formato = local["pasta"], local["formato"]
    if not pasta:
        return [], None
    nomes = []
    try:
        def visitar(diretorio):
            modificacao_mais_recente = None
            with os.scandir(diretorio) as entradas:
                for entrada in entradas:
                    if entrada.name.startswith("."):
                        continue  # .DS_Store, ._* e outros ocultos do macOS
                    if entrada.is_dir():
                        if formato == "pastas_por_paciente":
                            nome = planilha.nome_do_arquivo(entrada.name)
                            try:
                                modificacao = visitar(entrada.path)
                            except OSError:
                                if ilegiveis is not None:
                                    ilegiveis.append(entrada.name)
                                try:
                                    modificacao = datetime.fromtimestamp(entrada.stat().st_mtime)
                                except OSError:
                                    modificacao = None
                            if nome:
                                nomes.append({"nome": nome, "entrada": entrada.name,
                                              "modificado": modificacao.date().isoformat() if modificacao else None})
                            if modificacao and (modificacao_mais_recente is None or
                                                modificacao > modificacao_mais_recente):
                                modificacao_mais_recente = modificacao
                    elif entrada.is_file():
                        modificacao = datetime.fromtimestamp(entrada.stat().st_mtime)
                        if formato == "arquivos_soltos":
                            nome = planilha.nome_do_arquivo(entrada.name)
                            if nome:
                                nomes.append({"nome": nome, "entrada": entrada.name,
                                              "modificado": modificacao.date().isoformat(),
                                              "do_mes": (modificacao.year, modificacao.month) == (ano, mes)})
                        else:
                            nome = planilha.nome_do_arquivo(entrada.name)
                            if nome:
                                nomes.append({"nome": nome, "entrada": entrada.name,
                                              "modificado": modificacao.date().isoformat()})
                            if modificacao_mais_recente is None or modificacao > modificacao_mais_recente:
                                modificacao_mais_recente = modificacao
            return modificacao_mais_recente
        visitar(pasta)
    except OSError as erro:
        return [], f"Não foi possível listar scanners: {erro}"
    return nomes, None


def conferir_scanners(fichas, nomes, disponivel=True):
    """Marca fichas abertas e separa nomes que não conferem com nenhuma linha."""
    resultado = []
    for ficha in fichas:
        item = dict(ficha)
        if disponivel and ficha["estado"] in {"por_auditar", "com_equipe"}:
            item["scanner"] = ("escaneada_aguardando_auditoria" if any(
                nome.get("do_mes", True) and planilha.nomes_conferem(nome["nome"], ficha["paciente"])
                for nome in nomes)
                else "falta_escanear")
            if ficha["estado"] == "com_equipe" and ficha.get("ultimo_envio"):
                posteriores = [nome["modificado"] for nome in nomes
                               if nome.get("modificado") and
                               nome["modificado"] > ficha["ultimo_envio"]["data"] and
                               planilha.nomes_conferem(nome["nome"], ficha["paciente"])]
                if posteriores:
                    item["scanner"] = "arquivo_novo_apos_aviso"
                    item["arquivo_novo_em"] = max(posteriores)
        resultado.append(item)
    sem_linha = []
    vistos = set()
    for nome in nomes:
        if not nome.get("do_mes", True):
            continue
        if (not any(planilha.nomes_conferem(nome["nome"], ficha["paciente"]) for ficha in fichas)
                and nome["nome"] not in vistos):
            sem_linha.append(nome)
            vistos.add(nome["nome"])
    return resultado, sem_linha


def pode_avisar(ficha):
    """Novo aviso à equipe só sem aviso anterior ou quando chegou arquivo novo depois dele.

    Depois do aviso o prazo da equipe já está correndo; outro aviso só faz sentido quando a
    equipe reenviou e a revisão achou nova pendência. Sem leitura dos scanners não há como
    saber se a equipe reenviou, então o aviso continua liberado.
    """
    return ("ultimo_envio" not in ficha or ficha.get("scanner") is None
            or ficha["scanner"] == "arquivo_novo_apos_aviso")

def _cor_da_celula(celula):
    preenchimento = celula.fill
    if preenchimento.patternType is None:
        return None
    cor = preenchimento.fgColor
    if preenchimento.patternType != "solid" or cor.type != "rgb" or cor.tint != 0:
        return "cor_desconhecida"
    rgb = cor.rgb
    if not isinstance(rgb, str) or len(rgb) != 8:
        return "cor_desconhecida"
    rgb = rgb.upper()
    if rgb in VERMELHOS:
        return "com_equipe"
    if rgb in AMARELOS:
        return "por_auditar"
    return "cor_desconhecida"


def cor_da_linha(aba, linha, col_numero="B", col_observacoes="T"):
    """Vermelho prevalece; qualquer cor ilegível impede aprovação silenciosa."""
    inicio = column_index_from_string(col_numero)
    fim = column_index_from_string(col_observacoes)
    estados = {_cor_da_celula(aba.cell(linha, coluna)) for coluna in range(inicio, fim + 1)}
    if "com_equipe" in estados:
        return "com_equipe"
    if "cor_desconhecida" in estados:
        return "cor_desconhecida"
    if "por_auditar" in estados:
        return "por_auditar"
    return "concluida"


def comentario_da_linha(aba, linha, col_numero="B", col_observacoes="T"):
    """Primeiro comentário entre Nº e observações, sem o prefixo Autor:."""
    for coluna in range(column_index_from_string(col_numero),
                        column_index_from_string(col_observacoes) + 1):
        comentario = aba.cell(linha, coluna).comment
        if comentario:
            return re.sub(r"^\s*Autor:\s*", "", comentario.text.strip(), flags=re.IGNORECASE)
    return ""


def ler_fichas(caminho, abas):
    """Lê fichas das abas configuradas, parando na primeira linha sem paciente."""
    livro = load_workbook(caminho, data_only=True)
    fichas = []
    try:
        nomes = {planilha.normalizar(nome): nome for nome in livro.sheetnames}
        for configuracao in abas:
            nome_aba = next((nomes[planilha.normalizar(nome)] for nome in configuracao.nomes
                             if planilha.normalizar(nome) in nomes), None)
            if nome_aba is None:
                continue
            aba = livro[nome_aba]
            for linha in range(configuracao.primeira_linha, aba.max_row + 1):
                paciente = aba[f"{configuracao.col_nome}{linha}"].value
                if paciente is None or not str(paciente).strip():
                    break
                numero = planilha.numero_para_texto(aba[f"{configuracao.col_numero}{linha}"].value)
                posicao = linha - configuracao.primeira_linha + 1
                numero_codigo = int(numero) if numero.isdigit() else numero
                fichas.append({"ficha": f"{configuracao.codigo}{numero_codigo:02d}" if isinstance(numero_codigo, int)
                               else f"{configuracao.codigo}{numero}",
                               "numero": numero, "paciente": planilha.normalizar(paciente),
                               "tipo": configuracao.tipo, "aba": nome_aba, "linha": linha,
                               "posicao": posicao, "estado": cor_da_linha(aba, linha, configuracao.col_numero),
                               "comentario": comentario_da_linha(aba, linha, configuracao.col_numero)})
    finally:
        livro.close()
    return fichas


def listar_pdfs(pasta_mes, aba):
    """Lista PDFs numerados da subpasta configurada, sem criá-la."""
    if not pasta_mes or not Path(pasta_mes).is_dir():
        return []
    alvo = planilha.normalizar(aba.palavra_subpasta or aba.tipo.split("-")[0])
    pastas = sorted(p for p in Path(pasta_mes).iterdir() if p.is_dir()
                    and planilha.normalizar(p.name).startswith(alvo)
                    and ("COVID" in alvo or "COVID" not in planilha.normalizar(p.name)))
    if not pastas:
        return []
    encontrados = []
    for caminho in sorted(pastas[0].iterdir()):
        if not caminho.is_file():
            continue
        resultado = PADRAO_PDF.fullmatch(caminho.name)
        if resultado:
            encontrados.append({"numero": resultado.group(1), "paciente": resultado.group(2),
                                "caminho": caminho})
    return encontrados


def cruzar(fichas, pdfs_por_tipo):
    """Cruza fichas com PDFs e devolve fichas, inconsistências e PDFs sem linha."""
    resultado = []
    inconsistencias = []
    sem_linha = []
    por_tipo = {}
    for ficha in fichas:
        por_tipo.setdefault(ficha["tipo"], []).append(ficha)
    for tipo in dict.fromkeys([*por_tipo, *pdfs_por_tipo]):
        linhas = por_tipo.get(tipo, [])
        arquivos = pdfs_por_tipo.get(tipo, [])
        contagem = Counter(planilha.numero_para_texto(p["numero"]).lstrip("0") for p in arquivos)
        numeros = {planilha.numero_para_texto(f["numero"]).lstrip("0") for f in linhas}
        for arquivo in arquivos:
            numero = planilha.numero_para_texto(arquivo["numero"]).lstrip("0")
            if numero not in numeros:
                sem_linha.append({"tipo": tipo, **arquivo})
        for numero, quantidade in contagem.items():
            if quantidade > 1:
                inconsistencias.append({"tipo": tipo, "numero": numero, "motivo": "pdf_duplicado"})
        for ficha in linhas:
            numero = planilha.numero_para_texto(ficha["numero"]).lstrip("0")
            correspondentes = [p for p in arquivos if planilha.numero_para_texto(p["numero"]).lstrip("0") == numero]
            item = {**ficha, "tem_pdf": bool(correspondentes)}
            resultado.append(item)
            motivos = []
            if ficha["estado"] == "concluida" and not correspondentes:
                motivos.append("sem_cor_sem_pdf")
            if ficha["estado"] == "por_auditar" and correspondentes:
                motivos.append("amarela_com_pdf")
            if ficha["estado"] == "com_equipe" and correspondentes:
                motivos.append("vermelha_com_pdf")
            if ficha["estado"] == "concluida" and ficha["comentario"]:
                motivos.append("sem_cor_com_comentario")
            if ficha["estado"] == "cor_desconhecida":
                motivos.append("cor_desconhecida")
            if not planilha.mesmo_numero(ficha["numero"], ficha["posicao"]):
                motivos.append("numero_fora_da_posicao")
            for arquivo in correspondentes:
                nome_pdf = planilha.normalizar(pdfs.limpar_nome_arquivo(arquivo["paciente"]))
                nome_ficha = planilha.normalizar(pdfs.limpar_nome_arquivo(ficha["paciente"]))
                if nome_pdf != nome_ficha:
                    motivos.append("nome_pdf_diferente")
            inconsistencias.extend({"ficha": ficha["ficha"], "tipo": tipo, "motivo": motivo}
                                    for motivo in dict.fromkeys(motivos))
    return {"fichas": resultado, "inconsistencias": inconsistencias, "pdfs_sem_linha": sem_linha}


def contar(fichas):
    """Conta estados; a porcentagem concluída usa todas as linhas lançadas."""
    estados = Counter(ficha["estado"] for ficha in fichas)
    total = len(fichas)
    return {"lancadas": total, "por_auditar": estados["por_auditar"],
            "com_equipe": estados["com_equipe"], "concluidas": estados["concluida"],
            "cor_desconhecida": estados["cor_desconhecida"],
            "percentual_concluido": 100 * estados["concluida"] / total if total else 0}


def pendencias(fichas):
    """Fichas sem envio primeiro, seguidas das devolvidas há mais tempo."""
    return sorted((ficha for ficha in fichas if ficha["estado"] == "com_equipe"),
                  key=lambda ficha: ("dias" in ficha, -ficha.get("dias", 0)))
