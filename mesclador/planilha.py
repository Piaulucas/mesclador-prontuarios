"""Tudo que envolve a planilha de faturamento: listar, ler e buscar o Nº da remoção."""

import difflib
import re
import unicodedata
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string

from .config import MESES

_cache = {}   # (caminho, data de modificação) -> remoções lidas


def normalizar(texto):
    """'José  da Silva ' -> 'JOSE DA SILVA' (sem acento, maiúsculo, espaços únicos)."""
    texto = unicodedata.normalize("NFKD", str(texto))
    texto = "".join(c for c in texto if not unicodedata.combining(c))
    return re.sub(r"\s+", " ", texto).strip().upper()


def numero_para_texto(valor):
    """45, 45.0, '45' ou ' 045 ' -> '45' / '045' (só tira o '.0' e os espaços)."""
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    return str(valor).strip()


def mesmo_numero(a, b):
    """'045' e '45' são o mesmo Nº."""
    return numero_para_texto(a).lstrip("0") == numero_para_texto(b).lstrip("0")


def _pontuacao_mes(caminho_relativo, ano, mes):
    """Quanto o caminho do arquivo 'parece' ser do mês escolhido.
    2 = tem o mês e o ano, 1 = só o mês, 0 = nada."""
    texto = normalizar(caminho_relativo)
    tokens = set(re.split(r"[^A-Z0-9]+", texto))
    tem_mes = normalizar(MESES[mes - 1]) in tokens or f"{mes:02d}" in tokens or f"{ano}{mes:02d}" in tokens
    tem_ano = str(ano) in tokens or f"{ano}{mes:02d}" in tokens
    return (2 if tem_ano else 1) if tem_mes else 0


def listar_planilhas(base, ano=None, mes=None, limite=60, somente_do_mes=False):
    """Lista as planilhas (.xlsx/.xlsm) da pasta da empresa.
    As que batem com o mês/ano escolhidos vêm primeiro; depois, as mais recentes.
    Com somente_do_mes=True, devolve apenas as que batem com o mês/ano."""
    if not base or not Path(base).exists():
        return []
    base = Path(base)
    encontradas = []
    for p in base.rglob("*"):
        if p.suffix.lower() not in {".xlsx", ".xlsm"} or p.name.startswith("~$"):
            continue
        partes = p.relative_to(base).parts
        # ignora backups e versões antigas: pastas que começam com "_" e arquivos "BACKUP..."
        if any(parte.startswith("_") for parte in partes[:-1]) or p.name.upper().startswith("BACKUP"):
            continue
        relativo = str(p.relative_to(base))
        pontos = _pontuacao_mes(relativo, ano, mes) if ano and mes else 0
        encontradas.append((pontos, p.stat().st_mtime, relativo))
    if somente_do_mes:
        encontradas = [e for e in encontradas if e[0] == 2]
    encontradas.sort(key=lambda x: (x[0], x[1]), reverse=True)
    return [{"caminho": rel, "do_mes": pontos == 2} for pontos, _, rel in encontradas[:limite]]


def resolver_caminho(base, relativo):
    """Caminho relativo vindo da tela -> caminho real, sem permitir sair da pasta base."""
    if not base or not relativo:
        return None
    base = Path(base).resolve()
    caminho = (base / relativo).resolve()
    if base not in caminho.parents or not caminho.is_file():
        return None
    return caminho


def ler_planilha(caminho, abas):
    """Lê as abas configuradas e devolve [{'numero','nome','tipo','aba','linha'}].
    Usa cache enquanto o arquivo não for modificado."""
    caminho = Path(caminho)
    chave = (str(caminho), caminho.stat().st_mtime)
    if chave in _cache:
        return _cache[chave]

    wb = load_workbook(caminho, data_only=True, read_only=True)
    abas_da_planilha = {normalizar(n): n for n in wb.sheetnames}
    remocoes = []
    for aba in abas:
        # usa o primeiro nome configurado que existir nesta planilha
        nome_real = next((abas_da_planilha[normalizar(n)] for n in aba.nomes
                          if normalizar(n) in abas_da_planilha), None)
        if nome_real is None:
            continue
        ws = wb[nome_real]
        i_num = column_index_from_string(aba.col_numero) - 1
        i_nome = column_index_from_string(aba.col_nome) - 1
        ultima_col = max(i_num, i_nome) + 1
        for linha, valores in enumerate(ws.iter_rows(min_row=aba.primeira_linha, max_col=ultima_col,
                                                     values_only=True), start=aba.primeira_linha):
            valores = tuple(valores) + (None,) * (ultima_col - len(valores))
            numero, nome = valores[i_num], valores[i_nome]
            if numero is None or nome is None or str(nome).strip() == "":
                continue
            remocoes.append({"numero": numero_para_texto(numero), "nome": normalizar(nome),
                             "tipo": aba.tipo, "aba": nome_real, "linha": linha})
    wb.close()

    _cache.clear()           # guarda só a última planilha lida
    _cache[chave] = remocoes
    return remocoes


def buscar(remocoes, numero, tipo=None):
    """Remoções com esse Nº. A numeração recomeça em cada aba, então informe o tipo
    para buscar só na aba certa; sem tipo, devolve as de todas as abas."""
    if not numero_para_texto(numero):
        return []
    return [r for r in remocoes if mesmo_numero(r["numero"], numero) and (not tipo or r["tipo"] == tipo)]


# ---------------------------------------------------------------- conferência de nomes

_SUFIXO_DRIVE = re.compile(r"-\d{8}T\d{6}Z(-\d+)*$")   # "-20261014T190644Z-1-001"
_PARTICULAS = {"DE", "DA", "DO", "DAS", "DOS", "E"}


def nome_da_origem(texto):
    """Nome do paciente a partir do nome do zip/pasta, ou None se for genérico.
    'MARIA SOUZA-20261014T190644Z-001.zip' -> 'MARIA SOUZA'
    'drive-download-20261014T190644Z-1-001.zip' -> None"""
    if not texto:
        return None
    nome = re.sub(r"\.zip$", "", str(texto).strip(), flags=re.I)
    nome = _SUFIXO_DRIVE.sub("", nome)
    nome = normalizar(nome.replace("_", " "))
    if not nome or nome.startswith("DRIVE-DOWNLOAD") or not re.search(r"[A-Z]{2}", nome):
        return None
    return nome


def _tokens(nome):
    return [t for t in re.split(r"[^A-Z0-9]+", normalizar(nome)) if t and t not in _PARTICULAS]


def nomes_conferem(nome_origem, nome_planilha):
    """True se os nomes parecem ser da mesma pessoa. Aceita nome abreviado
    ('MARIA SOUZA' x 'MARIA FICTICIA DE SOUZA') e pequenas diferenças de digitação."""
    a, b = _tokens(nome_origem), _tokens(nome_planilha)
    if not a or not b:
        return False
    if a == b:
        return True
    curto, longo = (a, b) if len(a) <= len(b) else (b, a)
    # todas as palavras do nome curto aparecem no longo, e o primeiro nome coincide
    if curto[0] == longo[0] and all(t in longo for t in curto):
        return True
    return difflib.SequenceMatcher(None, " ".join(a), " ".join(b)).ratio() >= 0.88


def buscar_por_nome(remocoes, nome_origem):
    """Busca nas abas lidas, priorizando nomes idênticos e depois os mais parecidos."""
    if not nome_origem or not str(nome_origem).strip():
        return []
    tokens_origem = _tokens(nome_origem)
    nome_comparacao = " ".join(tokens_origem)
    achados = [r for r in remocoes if nomes_conferem(nome_origem, r["nome"])]
    return sorted(achados, key=lambda r: (
        tokens_origem != _tokens(r["nome"]),
        -difflib.SequenceMatcher(None, nome_comparacao, " ".join(_tokens(r["nome"]))).ratio(),
    ))
