"""Rotas do painel com planilhas e pastas inteiramente fictícias."""

import json
import os
from datetime import datetime

from openpyxl import load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import PatternFill

from mesclador.web import create_app
from mesclador import auditoria

from .conftest import criar_planilha


def test_tela_e_link_do_painel(cliente):
    resposta = cliente.get("/painel/")
    assert resposta.status_code == 200
    html = resposta.get_data(as_text=True)
    assert "<h1>Auditoria</h1>" in html and "Painel da auditoria" not in html
    assert 'id="painelMes"' in html
    assert 'type="month"' in html
    assert 'id="painelEmpresa"' in html
    assert 'id="painelRecarregar"' in html and "Ler de novo" in html
    assert 'href="/painel/"' in cliente.get("/").get_data(as_text=True)
    assert "textContent" in html
    assert "Avisei a equipe" in html
    assert "Parcial, reenviei" in html
    assert "if (!f.pode_avisar)" in html
    assert 'parametrosPainel.get("empresa")' in html
    assert "Equipe avisada em ${dataBR(f.ultimo_envio.data)} · aguardando o reenvio" in html
    assert "Nota do último aviso: " in html and "auditoria-nota" in html
    assert "Situação da correção" not in html
    assert 'parametrosPainel.get("foco")' in html
    assert 'f.dias > 10' in html
    assert 'f.scanner === "arquivo_novo_apos_aviso"' in html
    assert "no.textContent = String(texto)" in html
    # nenhuma data em ISO na tela: leitura e avisos passam por dataBR
    assert '"Leitura de " + dados.hoje' not in html
    assert "leitura de ${dataBR(dadosPainel.hoje)}" in html
    assert "Aviso registrado: ${ficha.ficha} em ${dataBR(alteracao.envio.data)}." in html


def test_painel_lista_as_empresas_configuradas_e_escolhe_pela_url(cliente):
    html = cliente.get("/painel/").get_data(as_text=True)
    assert 'const EMPRESAS_PAINEL = ["EMPRESA A", "EMPRESA B"];' in html
    # ?empresa= escolhe a inicial; sem ela (ou com uma desconhecida) abre "Todas"
    assert 'empresaDaUrl && EMPRESAS_PAINEL.includes(empresaDaUrl) ? empresaDaUrl : ""' in html
    assert '"Todas"' in html and '"Todas as empresas"' in html
    assert 'b.setAttribute("aria-pressed", String(valor === empresaAtual))' in html
    assert "anelProgresso(e.contagem, 34, false)" in html
    # trocar de empresa atualiza a URL sem recarregar e não mexe no contador de horas
    assert 'empresaAtual ? parametros.set("empresa", empresaAtual) : parametros.delete("empresa")' in html
    assert "history.replaceState" in html
    assert "/horas/auditoria" not in html


def test_painel_quadro_tem_colunas_por_situacao(cliente):
    html = cliente.get("/painel/").get_data(as_text=True)
    colunas = ['titulo: "Falta escanear"', 'titulo: "Reenviado · conferir"',
               'titulo: "Pode avisar"', 'titulo: "Aguardando reenvio"']
    posicoes = [html.index(coluna) for coluna in colunas]
    assert posicoes == sorted(posicoes)
    assert "Arquivo novo" not in html and "chegou arquivo novo" not in html
    assert "Reenviado em ${dataBR(f.arquivo_novo_em)} · conferir" in html
    # a coluna segue a regra do servidor: sem pode_avisar, só aguarda o reenvio
    assert 'if (f.scanner === "arquivo_novo_apos_aviso") return "reenviado";' in html
    assert 'if (!f.pode_avisar) return "aguardando";' in html
    assert 'f.estado === "por_auditar" && f.scanner === "falta_escanear" && filtro(f)' in html
    assert "Falta a equipe escanear" in html
    # cartão compacto com os dias na cor da faixa
    assert "auditoria-faixa-${f.faixa}" in html
    for faixa in ("amarelo", "laranja", "vermelho"):
        assert (f".auditoria-faixa-{faixa} {{ background: var(--faixa-{faixa}); "
                f"color: var(--faixa-{faixa}-txt); }}") in html


def test_painel_resume_colunas_quando_ha_muitas_pendencias(cliente):
    html = cliente.get("/painel/").get_data(as_text=True)
    assert "const MUITAS_PENDENCIAS = 15, MOSTRAR_POR_COLUNA = 5;" in html
    assert "const resumir = total >= MUITAS_PENDENCIAS;" in html
    assert "itens.slice(0, MOSTRAR_POR_COLUNA)" in html
    assert "`ver todas (${itens.length})`" in html
    # mais atrasadas primeiro: sem aviso registrado, depois mais dias com a equipe
    assert ('(("dias" in a.ficha) - ("dias" in b.ficha)) || '
            '((b.ficha.dias || 0) - (a.ficha.dias || 0))') in html


def test_painel_formulario_de_aviso_so_abre_no_cartao_clicado(cliente):
    html = cliente.get("/painel/").get_data(as_text=True)
    assert 'avisar.addEventListener("click", () => alternarFormulario(cartao, avisar, empresa, f))' in html
    assert "let formularioAberto = null;" in html
    assert "fecharFormulario();\n  if (mesmo) return;" in html
    assert 'avisar.setAttribute("aria-expanded", "false")' in html
    # nenhum formulário é montado junto com o cartão
    cartao = html[html.index("function cartaoPendencia"):html.index("function fecharFormulario")]
    assert 'elemento("form"' not in cartao
    assert "{envio: {data: data.value, pedido: nota.value}}" in html
    assert '{envio: {data: dadosPainel.hoje, pedido: "parcial"}}' in html
    assert 'evento.key === "Escape"' in html


def test_painel_rodape_sem_linha_e_inconsistencias(cliente):
    html = cliente.get("/painel/").get_data(as_text=True)
    assert 'listaRodape("Sem linha na planilha", semLinha' in html
    assert 'listaRodape("Inconsistências", inconsistencias' in html
    assert 'origem: "escaneado no Drive"' in html and 'origem: "PDF já mesclado"' in html
    assert 'numero_fora_da_posicao: "número fora da posição"' in html
    assert 'nome_pdf_diferente: "nome do PDF diferente"' in html
    assert 'id="painelRodape" class="auditoria-rodape"' in html
    assert "Expandir todas" not in html and 'elemento("details"' not in html


def test_painel_mantem_recortes_e_mes_compartilhado(cliente):
    html = cliente.get("/painel/").get_data(as_text=True)
    for chave in ("atrasadas", "arquivos-novos", "sem-escaneamento"):
        assert f'"{chave}"' in html
    assert 'focoPainel ? parametros.set("foco", focoPainel) : parametros.delete("foco")' in html
    assert 'id="painelRecortes"' in html and 'botao("", "Tudo")' in html
    assert 'lembrar("painelMes")' in html and 'guardar("painelMes", mes)' in html
    assert 'id="painelAviso"' in html and "Aviso registrado" in html


def test_dados_do_mes_cruzam_planilha_e_pdfs_sem_linha(config, pasta_planilhas, tmp_path):
    caminho = pasta_planilhas / "2026" / "10_Planilha_de_Faturamento_Exemplo_Outubro_26.xlsx"
    livro = load_workbook(caminho)
    aba = livro["Remoção ADULTO-PED"]
    aba["E16"].fill = PatternFill(fill_type="solid", fgColor="FFFFFF00")
    aba["E17"].fill = PatternFill(fill_type="solid", fgColor="FFFF0000")
    aba["T17"].comment = Comment("Autor: Enviar BDV fictício", "Autor Fictício")
    livro.save(caminho)
    livro.close()
    pasta = tmp_path / "prontuarios" / "2026" / "10_OUTUBRO"
    adulto = pasta / "ADULTO-PED"
    adulto.mkdir(parents=True)
    (adulto / "46 - ANA SOUZA.pdf").write_bytes(b"PDF ficticio")
    (adulto / "99 - PESSOA FICTICIA.pdf").write_bytes(b"PDF ficticio")
    config.auditoria_prontuarios = {"EMPRESA A": str(tmp_path / "prontuarios" / "{ano}" / "{mm}_{MES}")}
    cliente = create_app(config, pasta_temp=tmp_path / "temporarios").test_client()

    resposta = cliente.get("/painel/dados?mes=2026-10")
    assert resposta.status_code == 200
    dados = resposta.get_json()
    assert dados["mes"] == "2026-10" and "hoje" in dados
    assert dados["totais"]["lancadas"] == 4
    assert dados["totais"]["por_auditar"] == 1
    assert dados["totais"]["com_equipe"] == 1
    assert dados["totais"]["sem_devolucao"] == 1
    empresa = dados["empresas"][0]
    assert empresa["planilha"] == "2026/10_Planilha_de_Faturamento_Exemplo_Outubro_26.xlsx"
    assert empresa["contagem"]["percentual_concluido"] == 50
    assert empresa["pendencias"][0]["comentario"] == "Enviar BDV fictício"
    assert empresa["pendencias"][0]["devolucao"] == "devolução não registrada"
    assert "dias" not in empresa["pendencias"][0]
    assert empresa["pdfs_sem_linha"][0]["arquivo"] == "99 - PESSOA FICTICIA.pdf"
    assert any(i["motivo"] == "sem_cor_sem_pdf" for i in empresa["inconsistencias"])
    assert "Pasta de prontuários" in dados["empresas"][1]["aviso"]


def test_mes_invalido_e_anterior_ao_inicio(config, tmp_path):
    config.mes_inicial = "2026-10"
    cliente = create_app(config, pasta_temp=tmp_path / "temporarios").test_client()
    for mes in ("", "0000-01", "2026-13", "2026-9", "2026-09", "texto", "2026-10-extra"):
        assert cliente.get("/painel/dados", query_string={"mes": mes}).status_code == 400
    assert cliente.get("/painel/dados?mes=2026-10").status_code == 200
    assert 'min="2026-10"' in cliente.get("/painel/").get_data(as_text=True)


def test_planilha_com_defeito_em_uma_empresa_nao_derruba_outra(config, tmp_path):
    defeito = tmp_path / "defeito"
    defeito.mkdir()
    (defeito / "2026-10.xlsx").write_text("arquivo fictício ilegível", encoding="utf-8")
    boa = tmp_path / "boa"
    criar_planilha(boa / "2026-10.xlsx", [(1, "Ana Exemplo")], [])
    config.empresas = {"EMPRESA FICTÍCIA A": defeito, "EMPRESA FICTÍCIA B": boa}
    cliente = create_app(config, pasta_temp=tmp_path / "temporarios").test_client()
    resposta = cliente.get("/painel/dados?mes=2026-10")
    assert resposta.status_code == 200
    primeira, segunda = resposta.get_json()["empresas"]
    assert "Não foi possível ler a planilha" in primeira["aviso"]
    assert primeira["contagem"]["lancadas"] == 0
    assert segunda["contagem"]["lancadas"] == 1
    assert resposta.get_json()["totais"]["lancadas"] == 1


def test_registro_de_envio_correcao_e_contagem(config, pasta_planilhas, tmp_path):
    caminho = pasta_planilhas / "2026" / "10_Planilha_de_Faturamento_Exemplo_Outubro_26.xlsx"
    livro = load_workbook(caminho)
    livro["Remoção ADULTO-PED"]["E16"].fill = PatternFill(fill_type="solid", fgColor="FFFF0000")
    livro.save(caminho)
    livro.close()
    cliente = create_app(config, pasta_temp=tmp_path / "temporarios").test_client()
    base = {"mes": "2026-10", "empresa": "EMPRESA A", "ficha": "ADU45",
            "paciente": "José da Silva"}
    for data in ("2026-10-20", "2026-10-25"):
        resposta = cliente.post("/painel/registro", json={**base,
            "envio": {"data": data, "pedido": "Enviar BDV fictício"}})
        assert resposta.status_code == 200
    resposta = cliente.post("/painel/registro", json={**base,
        "correcao": {"situacao": "parcial", "data": "2026-10-28"}})
    assert resposta.status_code == 200
    arquivos = list(config.pasta_registros.glob("*.json"))
    assert len(arquivos) == 1
    registro = json.loads(arquivos[0].read_text(encoding="utf-8"))["fichas"]["ADU45"]
    assert registro["paciente"] == "JOSE DA SILVA"
    assert [e["data"] for e in registro["envios"]] == ["2026-10-20", "2026-10-25"]
    assert registro["correcao"] == {"situacao": "parcial", "data": "2026-10-28"}
    dados = cliente.get("/painel/dados?mes=2026-10").get_json()
    pendencia = dados["empresas"][0]["pendencias"][0]
    assert pendencia["ultimo_envio"]["data"] == "2026-10-25"
    assert pendencia["correcao"]["situacao"] == "parcial"
    assert dados["totais"]["sem_devolucao"] == 0
    assert dados["empresas"][0]["contagem"]["sem_devolucao"] == 0


def test_registro_invalido_nao_grava(config, pasta_planilhas, tmp_path):
    caminho = pasta_planilhas / "2026" / "10_Planilha_de_Faturamento_Exemplo_Outubro_26.xlsx"
    livro = load_workbook(caminho)
    livro["Remoção ADULTO-PED"]["E16"].fill = PatternFill(fill_type="solid", fgColor="FFFF0000")
    livro.save(caminho)
    livro.close()
    cliente = create_app(config, pasta_temp=tmp_path / "temporarios").test_client()
    base = {"mes": "2026-10", "empresa": "EMPRESA A", "ficha": "ADU45",
            "paciente": "José da Silva", "envio": {"data": "2026-10-20", "pedido": "BDV fictício"}}
    casos = [
        {**base, "empresa": "EMPRESA INEXISTENTE"},
        {**base, "empresa": ["EMPRESA A"]},
        {**base, "mes": "2026-13"},
        {**base, "ficha": "ADU99"},
        {**base, "paciente": "OUTRA PESSOA"},
        {**base, "envio": {"data": "2026-02-30", "pedido": "BDV fictício"}},
        {**base, "envio": {"data": "2026-10-20", "pedido": 123}},
        {**base, "correcao": {"situacao": "recebida", "data": "2026-10-28"}},
        {chave: valor for chave, valor in base.items() if chave != "envio"},
        {**base, "envio": None, "correcao": {"situacao": "invalida", "data": "2026-10-28"}},
    ]
    for caso in casos:
        assert cliente.post("/painel/registro", json=caso).status_code == 400
    assert not config.pasta_registros.exists()


def test_aviso_sem_nota_e_reenvio_parcial_reiniciam_deteccao(config, pasta_planilhas, tmp_path):
    caminho = pasta_planilhas / "2026" / "10_Planilha_de_Faturamento_Exemplo_Outubro_26.xlsx"
    livro = load_workbook(caminho)
    livro["Remoção ADULTO-PED"]["E16"].fill = PatternFill(fill_type="solid", fgColor="FFFF0000")
    livro.save(caminho)
    livro.close()
    raiz = tmp_path / "scanners_ficticios"
    raiz.mkdir()
    arquivo = raiz / "Jose da Silva (BDV).pdf"
    arquivo.touch()
    instante = datetime(2026, 10, 27, 12).timestamp()
    os.utime(arquivo, (instante, instante))
    config.auditoria_scanners = {"EMPRESA A": {"pasta": str(raiz), "formato": "arquivos_soltos"}}
    cliente = create_app(config, pasta_temp=tmp_path / "temporarios").test_client()
    base = {"mes": "2026-10", "empresa": "EMPRESA A", "ficha": "ADU45",
            "paciente": "José da Silva"}
    for nota in ({}, {"pedido": ""}):
        assert cliente.post("/painel/registro", json={**base,
            "envio": {"data": "2026-10-25", **nota}}).status_code == 200
    pendencia = cliente.get("/painel/dados?mes=2026-10").get_json()["empresas"][0]["pendencias"][0]
    assert pendencia["ultimo_envio"]["pedido"] == ""
    assert pendencia["scanner"] == "arquivo_novo_apos_aviso"
    assert pendencia["arquivo_novo_em"] == "2026-10-27"
    assert cliente.post("/painel/registro", json={**base,
        "envio": {"data": "2026-10-31", "pedido": "parcial"}}).status_code == 200
    pendencia = cliente.get("/painel/dados?mes=2026-10").get_json()["empresas"][0]["pendencias"][0]
    assert pendencia["ultimo_envio"] == {"data": "2026-10-31", "pedido": "parcial"}
    assert pendencia["scanner"] == "escaneada_aguardando_auditoria"
    assert "arquivo_novo_em" not in pendencia
    registro = auditoria.ler_registros(config, "EMPRESA A", "2026-10")["fichas"]["ADU45"]
    assert [envio["pedido"] for envio in registro["envios"]] == ["", "", "parcial"]


def test_painel_ignora_registro_de_outro_paciente(config, pasta_planilhas, tmp_path):
    caminho = pasta_planilhas / "2026" / "10_Planilha_de_Faturamento_Exemplo_Outubro_26.xlsx"
    livro = load_workbook(caminho)
    livro["Remoção ADULTO-PED"]["E16"].fill = PatternFill(fill_type="solid", fgColor="FFFF0000")
    livro.save(caminho)
    livro.close()
    auditoria.salvar_registros(config, "EMPRESA A", "2026-10", {"fichas": {
        "ADU45": {"paciente": "OUTRA PESSOA FICTICIA", "envios": [
            {"data": "2026-10-20", "pedido": "Pedido fictício"}]}}})
    cliente = create_app(config, pasta_temp=tmp_path / "temporarios").test_client()
    dados = cliente.get("/painel/dados?mes=2026-10").get_json()
    pendencia = dados["empresas"][0]["pendencias"][0]
    assert pendencia["devolucao"] == "devolução não registrada"
    assert "dias" not in pendencia
    assert dados["totais"]["sem_devolucao"] == 1
    assert {"ficha": "ADU45", "tipo": "ADULTO-PED", "motivo": "registro_outro_paciente"} in (
        dados["empresas"][0]["inconsistencias"])
    assert cliente.post("/painel/registro", json={"mes": "2026-10", "empresa": "EMPRESA A",
        "ficha": "ADU45", "paciente": "José da Silva",
        "envio": {"data": "2026-10-25", "pedido": "Novo pedido fictício"}}).status_code == 409


def test_painel_mostra_scanners_e_escaneados_sem_linha(config, pasta_planilhas, tmp_path):
    caminho = pasta_planilhas / "2026" / "10_Planilha_de_Faturamento_Exemplo_Outubro_26.xlsx"
    livro = load_workbook(caminho)
    livro["Remoção ADULTO-PED"]["E16"].fill = PatternFill(fill_type="solid", fgColor="FFFFFF00")
    livro["Remoção ADULTO-PED"]["E17"].fill = PatternFill(fill_type="solid", fgColor="FFFF0000")
    livro.save(caminho)
    livro.close()
    raiz = tmp_path / "scanners_ficticios"
    mes = raiz / "10-OUTUBRO-26"
    (mes / "Jose da Silva").mkdir(parents=True)
    (mes / "Pessoa Sem Linha").mkdir()
    config.auditoria_scanners = {"EMPRESA A": {"pasta": str(raiz), "formato": "pastas_por_paciente"}}
    cliente = create_app(config, pasta_temp=tmp_path / "temporarios").test_client()
    dados = cliente.get("/painel/dados?mes=2026-10").get_json()
    empresa = dados["empresas"][0]
    estados = {f["ficha"]: f.get("scanner") for f in empresa["fichas"]}
    assert estados["ADU45"] == "escaneada_aguardando_auditoria"
    assert estados["ADU03"] == "falta_escanear"
    assert estados["ADU46"] is None
    assert empresa["contagem"]["falta_escanear"] == 1
    assert dados["totais"]["falta_escanear"] == 1
    assert [n["entrada"] for n in empresa["escaneados_sem_linha"]] == ["Pessoa Sem Linha"]
    assert "Pasta de scanners" in dados["empresas"][1]["aviso"]
    html = cliente.get("/painel/").get_data(as_text=True)
    assert "Sem linha na planilha" in html
    assert "Falta a equipe escanear" in html


def test_mes_padrao_abre_o_anterior_ate_o_atual_ter_planilha(tmp_path):
    from datetime import date
    from types import SimpleNamespace
    from mesclador.painel import mes_padrao
    base = tmp_path / "empresa"
    (base / "2026").mkdir(parents=True)
    (base / "2026" / "OUTUBRO.xlsx").write_bytes(b"")
    config = SimpleNamespace(empresas={"A": base, "B": ""}, mes_inicial="2026-10")
    assert mes_padrao(config, date(2026, 11, 1)) == "2026-10"
    (base / "2026" / "NOVEMBRO.xlsx").write_bytes(b"")
    assert mes_padrao(config, date(2026, 11, 1)) == "2026-11"
    assert mes_padrao(config, date(2026, 10, 15)) == "2026-10"


def test_mes_padrao_respeita_mes_inicial_e_virada_de_ano(tmp_path):
    from datetime import date
    from types import SimpleNamespace
    from mesclador.painel import mes_padrao
    base = tmp_path / "empresa"
    (base / "2026").mkdir(parents=True)
    (base / "2026" / "DEZEMBRO.xlsx").write_bytes(b"")
    config = SimpleNamespace(empresas={"A": base}, mes_inicial="2026-10")
    assert mes_padrao(config, date(2027, 1, 3)) == "2026-12"
    config.mes_inicial = "2027-01"
    assert mes_padrao(config, date(2027, 1, 3)) == "2027-01"
    config.empresas = {"A": tmp_path / "nao_existe"}
    assert mes_padrao(config, date(2027, 1, 3)) == "2027-01"


def test_segundo_aviso_so_depois_de_arquivo_novo_da_equipe(config, pasta_planilhas, tmp_path):
    caminho = pasta_planilhas / "2026" / "10_Planilha_de_Faturamento_Exemplo_Outubro_26.xlsx"
    livro = load_workbook(caminho)
    livro["Remoção ADULTO-PED"]["E16"].fill = PatternFill(fill_type="solid", fgColor="FFFF0000")
    livro.save(caminho)
    livro.close()
    raiz = tmp_path / "scanners_ficticios"
    raiz.mkdir()
    arquivo = raiz / "Jose da Silva (BDV).pdf"
    arquivo.touch()
    os.utime(arquivo, (datetime(2026, 10, 10, 12).timestamp(),) * 2)
    config.auditoria_scanners = {"EMPRESA A": {"pasta": str(raiz), "formato": "arquivos_soltos"}}
    cliente = create_app(config, pasta_temp=tmp_path / "temporarios").test_client()
    base = {"mes": "2026-10", "empresa": "EMPRESA A", "ficha": "ADU45", "paciente": "José da Silva"}

    def pendencia():
        return cliente.get("/painel/dados?mes=2026-10").get_json()["empresas"][0]["pendencias"][0]

    assert pendencia()["pode_avisar"] is True
    assert cliente.post("/painel/registro", json={**base,
        "envio": {"data": "2026-10-20", "pedido": "BDV fictício"}}).status_code == 200
    assert pendencia()["pode_avisar"] is False
    resposta = cliente.post("/painel/registro", json={**base,
        "envio": {"data": "2026-10-22", "pedido": "de novo"}})
    assert resposta.status_code == 409
    assert "já foi avisada" in resposta.get_json()["erro"]
    # A equipe reenviou: arquivo novo depois do aviso libera um novo aviso.
    os.utime(arquivo, (datetime(2026, 10, 27, 12).timestamp(),) * 2)
    assert pendencia()["pode_avisar"] is True
    assert cliente.post("/painel/registro", json={**base,
        "envio": {"data": "2026-10-28", "pedido": "nova pendência"}}).status_code == 200
    assert pendencia()["pode_avisar"] is False
    registro = auditoria.ler_registros(config, "EMPRESA A", "2026-10")["fichas"]["ADU45"]
    assert [e["data"] for e in registro["envios"]] == ["2026-10-20", "2026-10-28"]


def test_pode_avisar():
    assert auditoria.pode_avisar({"scanner": "falta_escanear"})
    avisada = {"ultimo_envio": {"data": "2026-10-20", "pedido": ""}}
    assert auditoria.pode_avisar(avisada)  # sem leitura dos scanners
    for scanner, esperado in [("falta_escanear", False), ("escaneada_aguardando_auditoria", False),
                              ("arquivo_novo_apos_aviso", True)]:
        assert auditoria.pode_avisar({**avisada, "scanner": scanner}) is esperado
