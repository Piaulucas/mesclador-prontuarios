"""Auditoria com planilhas e nomes inteiramente fictícios."""

from datetime import date, datetime
import os

import pytest

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Color, PatternFill

from mesclador import auditoria
from mesclador.config import Aba, Config


ABAS = [Aba(["ADULTO-PED"], "ADULTO-PED", "B", "D", 16, "ADULTO"),
        Aba(["ADULTO-PED COVID"], "ADULTO-PED COVID", "B", "D", 16,
            "ADULTO_PED_COVID", "COV"),
        Aba(["NEONATAL COVID"], "NEONATAL COVID", "B", "D", 16,
            "NEONATAL_COVID", "NCV")]


def preencher(celula, cor):
    celula.fill = PatternFill(fill_type="solid", fgColor=cor)


def salvar_planilha(tmp_path, linhas):
    livro = Workbook()
    livro.active.title = ABAS[0].tipo
    for indice, configuracao in enumerate(ABAS):
        aba = livro.active if indice == 0 else livro.create_sheet(configuracao.tipo)
        for linha, numero, paciente, cor, comentario in linhas.get(configuracao.tipo, []):
            aba[f"B{linha}"] = numero
            aba[f"D{linha}"] = paciente
            if cor == "tema":
                aba[f"E{linha}"].fill = PatternFill(fill_type="solid", fgColor=Color(theme=1))
            elif cor == "indexada":
                aba[f"E{linha}"].fill = PatternFill(fill_type="solid", fgColor=Color(indexed=5))
            elif cor == "tint":
                aba[f"E{linha}"].fill = PatternFill(fill_type="solid", fgColor=Color(rgb="FFFFFF00", tint=0.2))
            elif cor:
                preencher(aba[f"E{linha}"], cor)
            if comentario:
                aba[f"T{linha}"].comment = Comment(comentario, "Autor Fictício")
    caminho = tmp_path / "faturamento_ficticio.xlsx"
    livro.save(caminho)
    return caminho


def criar_pdf(pasta, nome):
    pasta.mkdir(parents=True, exist_ok=True)
    (pasta / nome).write_bytes(b"PDF ficticio")


def test_cores_comentarios_total_e_contagem(tmp_path):
    linhas = {"ADULTO-PED": [
        (16, 1, "Ana Exemplo", "FFFFFF00", None),
        (17, 2, "Bruno Exemplo", "FFFF0000", "Autor: Enviar BDV"),
        (18, 3, "Clara Exemplo", None, None),
        (19, 4, "Dora Exemplo", "tema", None),
        (20, None, "   ", None, None),
        (21, 6, "Linha Após Total", None, None)]}
    fichas = auditoria.ler_fichas(salvar_planilha(tmp_path, linhas), ABAS)
    assert [f["estado"] for f in fichas] == [
        "por_auditar", "com_equipe", "concluida", "cor_desconhecida"]
    assert [f["ficha"] for f in fichas] == ["ADU01", "ADU02", "ADU03", "ADU04"]
    assert fichas[1]["comentario"] == "Enviar BDV"
    assert auditoria.contar(fichas) == {
        "lancadas": 4, "por_auditar": 1, "com_equipe": 1,
        "concluidas": 1, "cor_desconhecida": 1, "percentual_concluido": 25}
    assert [f["ficha"] for f in auditoria.pendencias(fichas)] == ["ADU02"]
    assert auditoria.contar([])["percentual_concluido"] == 0


def test_cor_vermelha_prevalece_e_cores_ilegiveis_nao_aprovam(tmp_path):
    caminho = salvar_planilha(tmp_path, {"ADULTO-PED": [
        (16, 1, "Ana Exemplo", "indexada", None),
        (17, 2, "Bruno Exemplo", "tint", None),
        (18, 3, "Clara Exemplo", "FF00FF00", None)]})
    from openpyxl import load_workbook
    livro = load_workbook(caminho)
    preencher(livro.active["F19"], "FFFFFF00")
    preencher(livro.active["G19"], "FFFFC7CE")
    livro.active["B19"] = 4
    livro.active["D19"] = "Dora Exemplo"
    livro.save(caminho)
    livro.close()
    assert [f["estado"] for f in auditoria.ler_fichas(caminho, ABAS)] == [
        "cor_desconhecida", "cor_desconhecida", "cor_desconhecida", "com_equipe"]


def test_cada_inconsistencia_e_pdfs_sem_linha(tmp_path):
    linhas = {"ADULTO-PED": [
        (16, 1, "Ana Exemplo", None, None),
        (17, 2, "Bruno Exemplo", "FFFFFF99", None),
        (18, 3, "Clara Exemplo", "FFFF9999", None),
        (19, 4, "Dora Exemplo", None, "Autor: Revisar"),
        (20, 8, "Eva Exemplo", "tema", None),
        (21, 6, "José Fictício", None, None)]}
    fichas = auditoria.ler_fichas(salvar_planilha(tmp_path, linhas), ABAS)
    pasta = tmp_path / "mes" / "ADULTO-PED"
    for nome in ["02 - Bruno Exemplo.pdf", "03 - Clara Exemplo.pdf",
                 "04 - Dora Exemplo.pdf", "06 - JOSE FICTICIO.pdf",
                 "06 - Outra Pessoa.pdf", "99 - Sem Linha.pdf"]:
        criar_pdf(pasta, nome)
    pdfs = auditoria.listar_pdfs(tmp_path / "mes", ABAS[0])
    dados = auditoria.cruzar(fichas, {"ADULTO-PED": pdfs})
    motivos = [i["motivo"] for i in dados["inconsistencias"]]
    assert set(motivos) == {
        "sem_cor_sem_pdf", "amarela_com_pdf", "vermelha_com_pdf",
        "sem_cor_com_comentario", "cor_desconhecida", "numero_fora_da_posicao",
        "pdf_duplicado", "nome_pdf_diferente"}
    assert motivos.count("nome_pdf_diferente") == 1
    assert [p["numero"] for p in dados["pdfs_sem_linha"]] == ["99"]
    assert dados["fichas"][0]["tem_pdf"] is False


def test_covid_tem_codigo_e_subpasta_proprios(tmp_path):
    caminho = salvar_planilha(tmp_path, {
        "ADULTO-PED": [(16, 1, "Ana Exemplo", None, None)],
        "ADULTO-PED COVID": [(16, 1, "Bia Exemplo", None, None)],
        "NEONATAL COVID": [(16, 1, "Cris Exemplo", None, None)]})
    fichas = auditoria.ler_fichas(caminho, ABAS)
    assert [f["ficha"] for f in fichas] == ["ADU01", "COV01", "NCV01"]
    criar_pdf(tmp_path / "mes" / "ADULTO_PED_COVID", "01 - Bia Exemplo.pdf")
    assert auditoria.listar_pdfs(tmp_path / "mes", ABAS[0]) == []
    assert len(auditoria.listar_pdfs(tmp_path / "mes", ABAS[1])) == 1
    assert len(auditoria.cruzar([], {"ADULTO-PED COVID": auditoria.listar_pdfs(
        tmp_path / "mes", ABAS[1])})["pdfs_sem_linha"]) == 1


def test_pasta_do_mes_ausente_e_variante_sem_acento(tmp_path):
    config = Config(tmp_path, {}, ABAS, auditoria_prontuarios={
        "EMPRESA FICTÍCIA": str(tmp_path / "{ano}" / "{mm}_{MES}")})
    assert auditoria.pasta_prontuarios(config, "EMPRESA FICTÍCIA", 2026, 3) is None
    assert auditoria.localizar_prontuarios(config, "EMPRESA FICTÍCIA", 2026, 3)["aviso"]
    assert auditoria.listar_pdfs(None, ABAS[0]) == []
    pasta = tmp_path / "2026" / "03_MARCO"
    pasta.mkdir(parents=True)
    assert auditoria.pasta_prontuarios(config, "EMPRESA FICTÍCIA", 2026, 3) == pasta
    assert auditoria.localizar_prontuarios(config, "EMPRESA FICTÍCIA", 2026, 3)["aviso"] is None


def test_ultimo_envio_define_prazo_e_registro_de_outro_paciente_e_ignorado(tmp_path):
    fichas = [
        {"ficha": "ADU01", "paciente": "ANA EXEMPLO", "tipo": "ADULTO-PED", "estado": "com_equipe"},
        {"ficha": "ADU02", "paciente": "BIA FICTICIA", "tipo": "ADULTO-PED", "estado": "com_equipe"},
        {"ficha": "ADU03", "paciente": "CAIO EXEMPLO", "tipo": "ADULTO-PED", "estado": "com_equipe"},
    ]
    registros = {"fichas": {
        "ADU01": {"paciente": "ANA EXEMPLO", "envios": [
            {"data": "2026-10-10", "pedido": "Primeiro pedido"},
            {"data": "2026-10-23", "pedido": "Segundo pedido"}]},
        "ADU02": {"paciente": "OUTRA PESSOA", "envios": [
            {"data": "2026-10-15", "pedido": "Pedido fictício"}]},
    }}
    resultado, inconsistencias = auditoria.aplicar_registros(fichas, registros, date(2026, 10, 30))
    assert resultado[0]["dias"] == 7
    assert resultado[0]["faixa"] == "laranja"
    assert resultado[0]["ultimo_envio"]["pedido"] == "Segundo pedido"
    assert "dias" not in resultado[1] and "dias" not in resultado[2]
    assert resultado[1]["devolucao"] == "devolução não registrada"
    assert inconsistencias == [{"ficha": "ADU02", "tipo": "ADULTO-PED",
                             "motivo": "registro_outro_paciente"}]
    assert [f["ficha"] for f in auditoria.pendencias(resultado)] == ["ADU02", "ADU03", "ADU01"]
    assert auditoria.aplicar_registros(fichas[:1], {"fichas": {"ADU01": {
        "paciente": "ANA EXEMPLO", "envios": [{"data": "2026-10-24", "pedido": "Pedido"}]}}},
        date(2026, 10, 30))[0][0]["faixa"] == "amarelo"
    assert auditoria.aplicar_registros(fichas[:1], {"fichas": {"ADU01": {
        "paciente": "ANA EXEMPLO", "envios": [{"data": "2026-10-19", "pedido": "Pedido"}]}}},
        date(2026, 10, 30))[0][0]["faixa"] == "vermelho"


def test_gravacao_atomica_preserva_arquivo_anterior_se_substituicao_falha(tmp_path, monkeypatch):
    config = Config(tmp_path, {}, ABAS, pasta_registros=tmp_path / "registros")
    anterior = {"fichas": {"ADU01": {"paciente": "ANA EXEMPLO", "envios": []}}}
    auditoria.salvar_registros(config, "EMPRESA FICTÍCIA", "2026-10", anterior)
    caminho = auditoria.caminho_registros(config, "EMPRESA FICTÍCIA", "2026-10")
    conteudo = caminho.read_bytes()

    def falhar_substituicao(*argumentos):
        raise OSError("Falha fictícia")

    monkeypatch.setattr(auditoria.os, "replace", falhar_substituicao)
    with pytest.raises(OSError):
        auditoria.salvar_registros(config, "EMPRESA FICTÍCIA", "2026-10", {"fichas": {}})
    assert caminho.read_bytes() == conteudo
    assert list(caminho.parent.iterdir()) == [caminho]


@pytest.mark.parametrize("nome_mes", ["10-OUTUBRO-26", "10 - OUTUBRO", "10 - FICHAS DE OUTUBRO"])
def test_scanners_em_pastas_com_nomes_variados(tmp_path, nome_mes):
    raiz = tmp_path / "scanners_ficticios"
    mes = raiz / nome_mes
    (mes / "Ana Exemplo" / "folhas").mkdir(parents=True)
    (mes / "Pessoa Sem Linha").mkdir()
    (raiz / "09 - SETEMBRO").mkdir()
    (raiz / "10-OUTUBRO-25").mkdir()
    config = Config(tmp_path, {}, ABAS, auditoria_scanners={
        "EMPRESA FICTÍCIA": {"pasta": str(raiz), "formato": "pastas_por_paciente"}})
    local = auditoria.localizar_scanners(config, "EMPRESA FICTÍCIA", 2026, 10)
    assert local["pasta"] == mes and local["aviso"] is None
    nomes, aviso = auditoria.listar_scanners(local, 2026, 10)
    assert aviso is None
    fichas, sem_linha = auditoria.conferir_scanners([
        {"paciente": "ANA EXEMPLO", "estado": "por_auditar"},
        {"paciente": "BRUNO EXEMPLO", "estado": "com_equipe"},
        {"paciente": "CLARA EXEMPLO", "estado": "concluida"}], nomes)
    assert [f.get("scanner") for f in fichas] == [
        "escaneada_aguardando_auditoria", "falta_escanear", None]
    assert [n["entrada"] for n in sem_linha] == ["Pessoa Sem Linha"]


def test_scanners_arquivos_soltos_filtram_mes_e_extraem_nome(tmp_path):
    raiz = tmp_path / "scanners_ficticios"
    raiz.mkdir()
    for nome, mes in [("Ana Exemplo (BDV).pdf", 10),
                      ("Pessoa Sem Linha (FORMULARIO).pdf", 10),
                      ("Bruno Exemplo (FDR).pdf", 9)]:
        caminho = raiz / nome
        caminho.touch()
        instante = datetime(2026, mes, 15).timestamp()
        os.utime(caminho, (instante, instante))
    config = Config(tmp_path, {}, ABAS, auditoria_scanners={
        "EMPRESA FICTÍCIA": {"pasta": str(raiz), "formato": "arquivos_soltos"}})
    local = auditoria.localizar_scanners(config, "EMPRESA FICTÍCIA", 2026, 10)
    nomes, aviso = auditoria.listar_scanners(local, 2026, 10)
    fichas, sem_linha = auditoria.conferir_scanners([
        {"paciente": "ANA EXEMPLO", "estado": "por_auditar"},
        {"paciente": "BRUNO EXEMPLO", "estado": "com_equipe"}], nomes)
    assert aviso is None
    assert [f["scanner"] for f in fichas] == ["escaneada_aguardando_auditoria", "falta_escanear"]
    assert [n["nome"] for n in sem_linha] == ["PESSOA SEM LINHA"]



@pytest.mark.parametrize("formato", ["arquivos_soltos", "pastas_por_paciente"])
def test_scanners_ignoram_arquivos_ocultos_do_macos(tmp_path, formato):
    (tmp_path / ".DS_Store").write_bytes(b"x")
    (tmp_path / "._Ana Exemplo (BDV).pdf").write_bytes(b"x")
    (tmp_path / ".pasta_oculta").mkdir()
    (tmp_path / "Ana Exemplo").mkdir()
    (tmp_path / "Ana Exemplo" / ".DS_Store").write_bytes(b"x")
    (tmp_path / "Ana Exemplo (BDV).pdf").touch()
    nomes, aviso = auditoria.listar_scanners({"pasta": tmp_path, "formato": formato}, 2026, 10)
    assert aviso is None
    assert all(not n["entrada"].startswith(".") for n in nomes)
    assert {n["nome"] for n in nomes} == {"ANA EXEMPLO"}

def test_scanners_em_pasta_usam_data_do_arquivo_mais_recente(tmp_path):
    pasta = tmp_path / "Ana Exemplo"
    (pasta / "folhas").mkdir(parents=True)
    for nome, dia in [("primeira.pdf", 20), ("folhas/segunda.pdf", 27)]:
        arquivo = pasta / nome
        arquivo.touch()
        instante = datetime(2026, 10, dia, 12).timestamp()
        os.utime(arquivo, (instante, instante))
    os.utime(pasta, (datetime(2026, 10, 30).timestamp(),) * 2)
    nomes, aviso = auditoria.listar_scanners(
        {"pasta": tmp_path, "formato": "pastas_por_paciente"}, 2026, 10)
    assert aviso is None
    assert next(nome for nome in nomes if nome["nome"] == "ANA EXEMPLO")["modificado"] == "2026-10-27"
    ficha = {"ficha": "ADU01", "paciente": "ANA EXEMPLO", "estado": "com_equipe",
             "ultimo_envio": {"data": "2026-10-25", "pedido": ""}}
    assert auditoria.conferir_scanners([ficha], nomes)[0][0]["arquivo_novo_em"] == "2026-10-27"
    for data in ("2026-10-27", "2026-10-28"):
        ficha["ultimo_envio"]["data"] = data
        assert auditoria.conferir_scanners([ficha], nomes)[0][0]["scanner"] == "escaneada_aguardando_auditoria"
    del ficha["ultimo_envio"]
    assert auditoria.conferir_scanners([ficha], nomes)[0][0]["scanner"] == "escaneada_aguardando_auditoria"


def test_scanners_soltos_detectam_arquivo_do_mes_seguinte_apos_aviso(tmp_path):
    arquivo = tmp_path / "Ana Exemplo (BDV).pdf"
    arquivo.touch()
    instante = datetime(2026, 11, 2, 12).timestamp()
    os.utime(arquivo, (instante, instante))
    nomes, aviso = auditoria.listar_scanners(
        {"pasta": tmp_path, "formato": "arquivos_soltos"}, 2026, 10)
    assert aviso is None
    assert nomes == [{"nome": "ANA EXEMPLO", "entrada": arquivo.name,
                      "modificado": "2026-11-02", "do_mes": False}]
    ficha = {"paciente": "ANA EXEMPLO", "estado": "com_equipe",
             "ultimo_envio": {"data": "2026-10-25", "pedido": ""}}
    conferidas, sem_linha = auditoria.conferir_scanners([ficha], nomes)
    assert conferidas[0]["scanner"] == "arquivo_novo_apos_aviso"
    assert conferidas[0]["arquivo_novo_em"] == "2026-11-02"
    assert sem_linha == []
    del ficha["ultimo_envio"]
    assert auditoria.conferir_scanners([ficha], nomes)[0][0]["scanner"] == "falta_escanear"


def test_scanners_pasta_ausente_e_duplicada(tmp_path):
    raiz = tmp_path / "scanners_ficticios"
    config = Config(tmp_path, {}, ABAS, auditoria_scanners={
        "EMPRESA FICTÍCIA": {"pasta": str(raiz), "formato": "pastas_por_paciente"}})
    assert "não encontrada" in auditoria.localizar_scanners(config, "EMPRESA FICTÍCIA", 2026, 10)["aviso"]
    raiz.mkdir()
    assert "do mês" in auditoria.localizar_scanners(config, "EMPRESA FICTÍCIA", 2026, 10)["aviso"]
    antiga = raiz / "10 - OUTUBRO"
    recente = raiz / "10 - FICHAS DE OUTUBRO"
    antiga.mkdir()
    recente.mkdir()
    os.utime(antiga, (100, 100))
    os.utime(recente, (200, 200))
    local = auditoria.localizar_scanners(config, "EMPRESA FICTÍCIA", 2026, 10)
    assert local["pasta"] == recente
    assert "Mais de uma" in local["aviso"]
