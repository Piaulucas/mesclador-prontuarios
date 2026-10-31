"""Dados fictícios para os testes. Nenhum dado real de paciente ou cliente."""

import io
import os
import time

import pytest
from openpyxl import Workbook
from PIL import Image
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from mesclador.config import Aba, Config
from mesclador.web import create_app

# Mesmo layout da planilha SESAB: cabeçalho na linha 15, Nº na coluna B, paciente na D
ABAS = [Aba(["ADULTO-PED", "Remoção ADULTO-PED"], "ADULTO-PED", "B", "D", 16, "ADULTO"),
        Aba(["NEONATAL", "Remoção NEONATAL"], "NEONATAL", "B", "D", 16, "NEO")]


def criar_planilha(caminho, adulto, neonatal, prefixo_aba=""):
    """adulto/neonatal: listas de (Nº, nome). prefixo_aba='Remoção ' imita a variação entre empresas."""
    wb = Workbook()
    for i, (tipo, linhas) in enumerate([("ADULTO-PED", adulto), ("NEONATAL", neonatal)]):
        ws = wb.active if i == 0 else wb.create_sheet()
        ws.title = prefixo_aba + tipo
        ws["B1"] = "PLANILHA DE FATURAMENTO (cabeçalho fictício)"
        for col, titulo in zip("BCDE", ["Nº DA REMOÇÃO", "DATA", "PACIENTE", "IDADE"]):
            ws[f"{col}15"] = titulo
        for n, (numero, nome) in enumerate(linhas, start=16):
            ws[f"B{n}"], ws[f"C{n}"], ws[f"D{n}"] = numero, "01/10/2026", nome
    caminho.parent.mkdir(parents=True, exist_ok=True)
    wb.save(caminho)


@pytest.fixture
def pasta_planilhas(tmp_path):
    base = tmp_path / "planilhas"
    criar_planilha(base / "2026" / "09_Planilha_de_Faturamento_Exemplo_Setembro_26.xlsx",
                   [(10, "Paciente Antigo")], [])
    time.sleep(0.01)
    criar_planilha(base / "2026" / "10_Planilha_de_Faturamento_Exemplo_Outubro_26.xlsx",
                   [(45, "José da Silva"), (3, "Maria Teste"), (46.0, "  ana   souza ")],
                   [(3, "RN de Fulana")], prefixo_aba="Remoção ")
    # versão antiga de outubro: deve ser ignorada
    criar_planilha(base / "2026" / "_versoes_anteriores" / "10_Planilha_de_Faturamento_Exemplo_Outubro_26_old.xlsx",
                   [(45, "NOME ANTIGO")], [])
    # a de setembro passa a ser a "mais recente" para provar que o mês tem prioridade
    os.utime(base / "2026" / "09_Planilha_de_Faturamento_Exemplo_Setembro_26.xlsx")
    return base


@pytest.fixture
def config(tmp_path, pasta_planilhas):
    return Config(pasta_inicial=tmp_path,
                  empresas={"EMPRESA A": pasta_planilhas, "EMPRESA B": None},
                  abas=ABAS, pasta_registros=tmp_path / "registros",
                  trabalho_arquivo=tmp_path / "trabalho.db")


@pytest.fixture
def cliente(config, tmp_path):
    app = create_app(config, pasta_temp=tmp_path / "temp")
    return app.test_client()


def jpeg(largura=600, altura=800):
    b = io.BytesIO()
    Image.new("RGB", (largura, altura), "white").save(b, "JPEG")
    b.seek(0)
    return b


def pdf(paginas=2):
    b = io.BytesIO()
    w = PdfWriter()
    for _ in range(paginas):
        w.add_blank_page(595, 842)
    w.write(b)
    b.seek(0)
    return b


def pdf_com_textos(textos):
    """PDF com uma página por texto (texto de verdade, extraível), para distinguir as páginas pelo conteúdo."""
    w = PdfWriter()
    for texto in textos:
        pagina = w.add_blank_page(595, 842)
        fonte = DictionaryObject({NameObject("/Type"): NameObject("/Font"), NameObject("/Subtype"): NameObject("/Type1"),
                                  NameObject("/BaseFont"): NameObject("/Helvetica")})
        pagina[NameObject("/Resources")] = DictionaryObject(
            {NameObject("/Font"): DictionaryObject({NameObject("/F1"): fonte})})
        conteudo = DecodedStreamObject()
        conteudo.set_data(f"BT /F1 48 Tf 100 700 Td ({texto}) Tj ET".encode("latin-1"))
        pagina[NameObject("/Contents")] = w._add_object(conteudo)
    b = io.BytesIO()
    w.write(b)
    b.seek(0)
    return b


def jpeg_colorido(cor, largura=600, altura=800, marca=None):
    """JPEG de uma cor só; `marca` pinta um quadrado (cor, tamanho) no canto superior esquerdo,
    para dar para saber se a imagem foi girada."""
    img = Image.new("RGB", (largura, altura), cor)
    if marca:
        quadrado = Image.new("RGB", (marca[1], marca[1]), marca[0])
        img.paste(quadrado, (0, 0))
    b = io.BytesIO()
    img.save(b, "JPEG", quality=95)
    b.seek(0)
    return b


@pytest.fixture
def pasta_principal(tmp_path):
    """Pasta do mês escolhida pelo usuário: já tem a subpasta de adulto, mas não a de neo."""
    p = tmp_path / "PRONTUARIOS OUTUBRO"
    (p / "Adulto e Pediátrico").mkdir(parents=True)
    return p
